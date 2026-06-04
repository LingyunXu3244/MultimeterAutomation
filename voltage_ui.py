import csv
import threading
import time
import tkinter as tk
from collections import deque
from datetime import datetime
from statistics import mean
from tkinter import filedialog, messagebox, ttk

import matplotlib
import pyvisa
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg
from matplotlib.figure import Figure

matplotlib.use("TkAgg")

RESOURCE = "GPIB0::6::INSTR"
MAX_POINTS = 300
MAX_RETRIES = 3

MODE_CONFIG = {
    "DC Voltage": {
        "conf": "CONF:VOLT:DC",
        "unit": "V",
        "plot": "DC Voltage",
        "supports_nplc": True,
        "supports_autozero": True,
        "supports_range": True,
        "range_cmd": "VOLT:DC:RANG",
        "nplc_cmd": "VOLT:DC:NPLC",
        "az_cmd": "VOLT:DC:ZERO:AUTO",
    },
    "AC Voltage": {
        "conf": "CONF:VOLT:AC",
        "unit": "V",
        "plot": "AC Voltage",
        "supports_nplc": False,
        "supports_autozero": False,
        "supports_range": True,
        "supports_aperture": False,
        "range_cmd": "VOLT:AC:RANG",
    },
    "Resistance": {
        "conf": "CONF:RES",
        "unit": "Ohm",
        "plot": "Resistance",
        "supports_nplc": True,
        "supports_autozero": True,
        "supports_range": True,
        "supports_aperture": False,
        "range_cmd": "RES:RANG",
        "nplc_cmd": "RES:NPLC",
        "az_cmd": "RES:ZERO:AUTO",
    },
    "Continuity": {
        "conf": "CONF:CONT",
        "unit": "Ohm",
        "plot": "Continuity",
        "supports_nplc": False,
        "supports_autozero": False,
        "supports_range": False,
        "supports_aperture": False,
    },
    "Diode": {
        "conf": "CONF:DIODE",
        "unit": "Vf",
        "plot": "Diode Forward Voltage",
        "supports_nplc": False,
        "supports_autozero": False,
        "supports_range": False,
        "supports_aperture": False,
    },
    "Frequency": {
        "conf": "CONF:FREQ",
        "unit": "Hz",
        "plot": "Frequency",
        "supports_nplc": False,
        "supports_autozero": False,
        "supports_range": True,
        "supports_aperture": True,
        "range_cmd": "FREQ:VOLT:RANG",
        "aperture_cmd": "FREQ:APER",
    },
    "Period": {
        "conf": "CONF:PER",
        "unit": "s",
        "plot": "Period",
        "supports_nplc": False,
        "supports_autozero": False,
        "supports_range": True,
        "supports_aperture": True,
        "range_cmd": "PER:VOLT:RANG",
        "aperture_cmd": "PER:APER",
    },
}

PRESETS = {
    "Fast": {"interval": 0.2, "nplc": 0.02, "autozero": False, "auto_range": True},
    "Stable Precision": {"interval": 1.0, "nplc": 10, "autozero": True, "auto_range": False},
    "Diode Quick": {"interval": 0.3, "nplc": 1, "autozero": False, "auto_range": True, "mode": "Diode"},
    "AC Monitor": {"interval": 0.5, "nplc": 1, "autozero": False, "auto_range": True, "mode": "AC Voltage"},
    "Frequency Gate": {"interval": 0.5, "nplc": 1, "autozero": False, "auto_range": True, "mode": "Frequency", "aperture": 0.1},
}


class VoltmeterApp:
    def __init__(self, root: tk.Tk) -> None:
        self.root = root
        self.root.title("HP 34401A - Multi-Mode Meter UI")
        self.root.geometry("1100x760")
        self.root.configure(bg="#1e1e1e")

        self.running = False
        self.dmm = None
        self.rm = None
        self._configured = False
        self._last_config_signature = None
        self._switching_mode = False

        self.readings = []
        self.chart_times = deque(maxlen=MAX_POINTS)
        self.chart_values = deque(maxlen=MAX_POINTS)
        self._start_time = None

        self._build_ui()
        self._build_chart()
        self._on_mode_changed()

    def _build_ui(self) -> None:
        top = tk.Frame(self.root, bg="#1e1e1e")
        top.pack(fill="x", padx=14, pady=(12, 0))

        tk.Label(
            top,
            text="HP 34401A Measurement Console",
            font=("Segoe UI", 14, "bold"),
            bg="#1e1e1e",
            fg="#c4c4c4",
        ).pack(side="left")

        self.status_var = tk.StringVar(value="Not connected")
        tk.Label(
            top,
            textvariable=self.status_var,
            font=("Segoe UI", 10),
            bg="#1e1e1e",
            fg="#7a7a7a",
        ).pack(side="right")

        display = tk.Frame(self.root, bg="#1e1e1e")
        display.pack(fill="x", padx=14, pady=(4, 0))

        self.value_var = tk.StringVar(value="---")
        self.value_label = tk.Label(
            display,
            textvariable=self.value_var,
            font=("Consolas", 56, "bold"),
            bg="#1e1e1e",
            fg="#00ff88",
        )
        self.value_label.pack(fill="x")

        stats_row = tk.Frame(display, bg="#1e1e1e")
        stats_row.pack(fill="x")
        self.count_var = tk.StringVar(value="Count: 0")
        self.min_var = tk.StringVar(value="Min: ---")
        self.max_var = tk.StringVar(value="Max: ---")
        self.mean_var = tk.StringVar(value="Mean: ---")
        for text_var in (self.count_var, self.min_var, self.max_var, self.mean_var):
            tk.Label(
                stats_row,
                textvariable=text_var,
                font=("Segoe UI", 9),
                bg="#1e1e1e",
                fg="#7f7f7f",
                padx=10,
            ).pack(side="left")

        controls = tk.LabelFrame(
            self.root,
            text="Controls",
            font=("Segoe UI", 9, "bold"),
            bg="#1e1e1e",
            fg="#9f9f9f",
            bd=1,
            relief="groove",
            padx=10,
            pady=8,
        )
        controls.pack(fill="x", padx=14, pady=8)

        tk.Label(controls, text="Mode", bg="#1e1e1e", fg="#bbbbbb").grid(row=0, column=0, sticky="w")
        self.mode_var = tk.StringVar(value="DC Voltage")
        self.mode_combo = ttk.Combobox(
            controls,
            textvariable=self.mode_var,
            values=list(MODE_CONFIG.keys()),
            state="readonly",
            width=14,
        )
        self.mode_combo.grid(row=1, column=0, padx=(0, 10), pady=(2, 6), sticky="w")
        self.mode_combo.bind("<<ComboboxSelected>>", lambda _e: self._on_mode_changed())

        tk.Label(controls, text="Interval (s)", bg="#1e1e1e", fg="#bbbbbb").grid(row=0, column=1, sticky="w")
        self.interval_var = tk.StringVar(value="0.5")
        tk.Entry(controls, textvariable=self.interval_var, width=8).grid(row=1, column=1, padx=(0, 10), sticky="w")

        tk.Label(controls, text="NPLC", bg="#1e1e1e", fg="#bbbbbb").grid(row=0, column=2, sticky="w")
        self.nplc_var = tk.StringVar(value="1")
        self.nplc_entry = tk.Entry(controls, textvariable=self.nplc_var, width=8)
        self.nplc_entry.grid(row=1, column=2, padx=(0, 10), sticky="w")

        self.autozero_var = tk.BooleanVar(value=True)
        self.autozero_check = tk.Checkbutton(
            controls,
            text="Autozero",
            variable=self.autozero_var,
            bg="#1e1e1e",
            fg="#bbbbbb",
            selectcolor="#1e1e1e",
            activebackground="#1e1e1e",
        )
        self.autozero_check.grid(row=1, column=3, padx=(0, 10), sticky="w")

        self.auto_range_var = tk.BooleanVar(value=True)
        self.auto_range_check = tk.Checkbutton(
            controls,
            text="Auto range",
            variable=self.auto_range_var,
            bg="#1e1e1e",
            fg="#bbbbbb",
            selectcolor="#1e1e1e",
            activebackground="#1e1e1e",
            command=self._toggle_range_entry,
        )
        self.auto_range_check.grid(row=1, column=4, padx=(0, 10), sticky="w")

        tk.Label(controls, text="Range", bg="#1e1e1e", fg="#bbbbbb").grid(row=0, column=5, sticky="w")
        self.range_var = tk.StringVar(value="10")
        self.range_entry = tk.Entry(controls, textvariable=self.range_var, width=8)
        self.range_entry.grid(row=1, column=5, padx=(0, 10), sticky="w")

        tk.Label(controls, text="Aperture (s)", bg="#1e1e1e", fg="#bbbbbb").grid(row=0, column=6, sticky="w")
        self.aperture_var = tk.StringVar(value="0.1")
        self.aperture_entry = tk.Entry(controls, textvariable=self.aperture_var, width=9)
        self.aperture_entry.grid(row=1, column=6, padx=(0, 10), sticky="w")

        tk.Label(controls, text="Low alarm", bg="#1e1e1e", fg="#bbbbbb").grid(row=0, column=7, sticky="w")
        self.low_alarm_var = tk.StringVar(value="")
        tk.Entry(controls, textvariable=self.low_alarm_var, width=9).grid(row=1, column=7, padx=(0, 6), sticky="w")

        tk.Label(controls, text="High alarm", bg="#1e1e1e", fg="#bbbbbb").grid(row=0, column=8, sticky="w")
        self.high_alarm_var = tk.StringVar(value="")
        tk.Entry(controls, textvariable=self.high_alarm_var, width=9).grid(row=1, column=8, padx=(0, 6), sticky="w")

        self.alarm_enabled_var = tk.BooleanVar(value=False)
        tk.Checkbutton(
            controls,
            text="Enable alarm",
            variable=self.alarm_enabled_var,
            bg="#1e1e1e",
            fg="#bbbbbb",
            selectcolor="#1e1e1e",
            activebackground="#1e1e1e",
        ).grid(row=1, column=9, padx=(0, 8), sticky="w")

        tk.Label(controls, text="Preset", bg="#1e1e1e", fg="#bbbbbb").grid(row=2, column=0, sticky="w")
        self.preset_var = tk.StringVar(value="Fast")
        ttk.Combobox(
            controls,
            textvariable=self.preset_var,
            values=list(PRESETS.keys()),
            state="readonly",
            width=14,
        ).grid(row=3, column=0, padx=(0, 10), pady=(2, 0), sticky="w")

        tk.Button(
            controls,
            text="Apply Preset",
            bg="#395d9b",
            fg="white",
            relief="flat",
            command=self.apply_preset,
            padx=10,
        ).grid(row=3, column=1, sticky="w")

        tk.Button(
            controls,
            text="Single Read",
            bg="#3b3b3b",
            fg="white",
            relief="flat",
            command=self.single_read,
            padx=10,
        ).grid(row=3, column=2, sticky="w")

        actions = tk.Frame(self.root, bg="#1e1e1e")
        actions.pack(fill="x", padx=14, pady=(2, 8))

        btn_style = {"font": ("Segoe UI", 10, "bold"), "relief": "flat", "padx": 12, "pady": 5}
        self.start_btn = tk.Button(actions, text="Connect & Start", bg="#007acc", fg="white", command=self.start, **btn_style)
        self.start_btn.pack(side="left", padx=4)
        self.stop_btn = tk.Button(actions, text="Stop", bg="#cc3333", fg="white", command=self.stop, state="disabled", **btn_style)
        self.stop_btn.pack(side="left", padx=4)
        tk.Button(actions, text="Export CSV", bg="#444444", fg="white", command=self.export_csv, **btn_style).pack(side="left", padx=4)
        tk.Button(actions, text="Export Excel", bg="#217346", fg="white", command=self.export_excel, **btn_style).pack(side="left", padx=4)
        tk.Button(actions, text="Clear Data", bg="#555555", fg="white", command=self.clear_data, **btn_style).pack(side="left", padx=4)

        self.root.protocol("WM_DELETE_WINDOW", self.on_close)

    def _build_chart(self) -> None:
        frame = tk.Frame(self.root, bg="#1e1e1e")
        frame.pack(fill="both", expand=True, padx=14, pady=(0, 12))

        self.fig = Figure(figsize=(10, 3.2), dpi=96, facecolor="#1e1e1e")
        self.ax = self.fig.add_subplot(111)
        self._style_axes()
        self.line, = self.ax.plot([], [], color="#00ff88", linewidth=1.4)

        self.canvas = FigureCanvasTkAgg(self.fig, master=frame)
        self.canvas.get_tk_widget().pack(fill="both", expand=True)
        self._schedule_chart_update()

    def _style_axes(self) -> None:
        unit = MODE_CONFIG[self.mode_var.get()]["unit"]
        plot_name = MODE_CONFIG[self.mode_var.get()]["plot"]
        self.ax.set_facecolor("#121212")
        self.ax.tick_params(colors="#888888", labelsize=8)
        self.ax.set_xlabel("Time (s)", color="#888888", fontsize=9)
        self.ax.set_ylabel(f"{plot_name} ({unit})", color="#888888", fontsize=9)
        self.ax.set_title(f"{plot_name} over Time", color="#aaaaaa", fontsize=10)
        for spine in self.ax.spines.values():
            spine.set_edgecolor("#333333")
        self.fig.tight_layout(pad=1.2)

    def _toggle_range_entry(self) -> None:
        state = "disabled" if self.auto_range_var.get() else "normal"
        self.range_entry.config(state=state)
        self._configured = False

    def _on_mode_changed(self) -> None:
        mode = MODE_CONFIG[self.mode_var.get()]
        self.nplc_entry.config(state="normal" if mode["supports_nplc"] else "disabled")
        self.autozero_check.config(state="normal" if mode["supports_autozero"] else "disabled")
        self.auto_range_check.config(state="normal" if mode["supports_range"] else "disabled")
        self.range_entry.config(state="normal" if mode["supports_range"] and not self.auto_range_var.get() else "disabled")
        self.aperture_entry.config(state="normal" if mode.get("supports_aperture") else "disabled")
        self.value_var.set(f"--- {mode['unit']}")
        self._style_axes()
        self.clear_data(reset_display=False)
        self.canvas.draw_idle()
        self._configured = False

        if self.running and not self._switching_mode:
            self._restart_for_mode_change()

    def _schedule_chart_update(self) -> None:
        if self.chart_times:
            self.line.set_data(list(self.chart_times), list(self.chart_values))
            self.ax.relim()
            self.ax.autoscale_view()
            self.canvas.draw_idle()
        self.root.after(500, self._schedule_chart_update)

    def _parse_reading(self, raw: str):
        text = raw.strip()
        if text.upper().startswith("OL"):
            return None
        return float(text)

    def _format_value(self, value):
        unit = MODE_CONFIG[self.mode_var.get()]["unit"]
        if value is None:
            return f"OL {unit}"
        if abs(value) >= 1000 or (abs(value) < 1e-3 and value != 0):
            return f"{value:+.6E} {unit}"
        return f"{value:+.6f} {unit}"

    def _update_stats(self) -> None:
        values = [r[2] for r in self.readings if r[2] is not None and r[0] == self.mode_var.get()]
        self.count_var.set(f"Count: {len(values)}")
        if not values:
            self.min_var.set("Min: ---")
            self.max_var.set("Max: ---")
            self.mean_var.set("Mean: ---")
            return
        unit = MODE_CONFIG[self.mode_var.get()]["unit"]
        self.min_var.set(f"Min: {min(values):.6g} {unit}")
        self.max_var.set(f"Max: {max(values):.6g} {unit}")
        self.mean_var.set(f"Mean: {mean(values):.6g} {unit}")

    def _apply_alarm_style(self, value):
        if value is None:
            self.value_label.config(fg="#ffcc66")
            return
        if not self.alarm_enabled_var.get():
            self.value_label.config(fg="#00ff88")
            return

        low = self._safe_float(self.low_alarm_var.get())
        high = self._safe_float(self.high_alarm_var.get())
        in_alarm = False
        if low is not None and value < low:
            in_alarm = True
        if high is not None and value > high:
            in_alarm = True
        self.value_label.config(fg="#ff4444" if in_alarm else "#00ff88")

    @staticmethod
    def _safe_float(text):
        if text is None or text.strip() == "":
            return None
        try:
            return float(text)
        except ValueError:
            return None

    def _configure_measurement_mode(self):
        if self.dmm is None:
            raise RuntimeError("Not connected to the multimeter")

        mode_name = self.mode_var.get()
        mode = MODE_CONFIG[mode_name]

        # Clear any pending errors / queues so the ERR annunciator does not
        # stay lit from a previous session, and make sure the buffer is empty
        # before sending the new CONF command.
        try:
            self.dmm.clear()
        except Exception:
            pass
        self.dmm.write("*CLS")

        # Send the function-change command and wait for the meter to finish
        # switching function before issuing any function-specific sub-commands
        # (RES/FREQ in particular need this; otherwise the sub-commands can be
        # parsed against the previous function and the front-panel display
        # does not switch).
        self.dmm.write(mode["conf"])
        try:
            self.dmm.query("*OPC?")
        except Exception:
            pass
        time.sleep(0.1)

        if mode["supports_range"]:
            # HP 34401A syntax: "<func>:RANG:AUTO ON|OFF" for autorange,
            # "<func>:RANG <value>" for a fixed range. Sending
            # "<func>:RANG AUTO" is invalid and lights the ERR annunciator.
            if self.auto_range_var.get():
                self.dmm.write(f"{mode['range_cmd']}:AUTO ON")
            else:
                self.dmm.write(f"{mode['range_cmd']}:AUTO OFF")
                rng = self._safe_float(self.range_var.get())
                if rng is not None:
                    self.dmm.write(f"{mode['range_cmd']} {rng}")

        if mode["supports_nplc"]:
            nplc = self._safe_float(self.nplc_var.get())
            if nplc is not None:
                self.dmm.write(f"{mode['nplc_cmd']} {nplc}")

        if mode["supports_autozero"]:
            self.dmm.write(f"{mode['az_cmd']} {'ON' if self.autozero_var.get() else 'OFF'}")

        if mode.get("supports_aperture"):
            aperture = self._safe_float(self.aperture_var.get())
            if aperture is not None:
                self.dmm.write(f"{mode['aperture_cmd']} {aperture}")

        self._configured = True

    def apply_preset(self) -> None:
        preset = PRESETS[self.preset_var.get()]
        self.interval_var.set(str(preset["interval"]))
        self.nplc_var.set(str(preset["nplc"]))
        self.autozero_var.set(bool(preset["autozero"]))
        self.auto_range_var.set(bool(preset["auto_range"]))
        if "aperture" in preset:
            self.aperture_var.set(str(preset["aperture"]))
        if "mode" in preset:
            self.mode_var.set(preset["mode"])
        self._on_mode_changed()
        self._toggle_range_entry()

    def single_read(self) -> None:
        if self.running:
            messagebox.showinfo("Busy", "Stop continuous mode before single read.")
            return
        rm = None
        dmm = None
        try:
            rm = pyvisa.ResourceManager("@ivi")
            dmm = rm.open_resource(RESOURCE)
            dmm.timeout = 5000
            self.dmm = dmm
            self._configure_measurement_mode()
            raw = dmm.query("READ?")
            value = self._parse_reading(raw)
            self.value_var.set(self._format_value(value))
            self._apply_alarm_style(value)
            self.status_var.set("Single read complete")
            now = datetime.now().strftime("%Y-%m-%d %H:%M:%S.%f")[:-3]
            settings = self._settings_string()
            self.readings.append((self.mode_var.get(), now, value, settings))
            self._update_stats()
        except Exception as exc:
            self.status_var.set(f"Error: {exc}")
            self.value_var.set("ERROR")
        finally:
            if dmm is not None:
                try:
                    dmm.close()
                except Exception:
                    pass
            if rm is not None:
                try:
                    rm.close()
                except Exception:
                    pass
            self.dmm = None

    def start(self) -> None:
        # If a previous read-loop thread is still winding down, wait for it
        # to finish so its `finally` block cannot null out the new session.
        prev = getattr(self, "_read_thread", None)
        if prev is not None and prev.is_alive():
            self.running = False
            prev.join(timeout=2.0)
        self.running = True
        self._start_time = None
        self.start_btn.config(state="disabled")
        self.stop_btn.config(state="normal")
        self.status_var.set("Connecting...")
        self._read_thread = threading.Thread(target=self._read_loop, daemon=True)
        self._read_thread.start()

    def stop(self) -> None:
        self.running = False
        self.start_btn.config(state="normal")
        self.stop_btn.config(state="disabled")
        self.status_var.set("Stopped.")

    def _restart_for_mode_change(self) -> None:
        self._switching_mode = True
        self.running = False
        self.status_var.set(f"Switching to {self.mode_var.get()}...")
        self.start_btn.config(state="disabled")
        self.stop_btn.config(state="disabled")
        # Wait long enough for the current read-loop iteration plus VISA
        # close to complete before respawning. `start()` will also join
        # the old thread defensively.
        self.root.after(700, self._finish_mode_restart)

    def _finish_mode_restart(self) -> None:
        self._configured = False
        self._last_config_signature = None
        self._switching_mode = False
        self.start()

    def _settings_string(self):
        return (
            f"interval={self.interval_var.get()},"
            f"nplc={self.nplc_var.get()},"
            f"aperture={self.aperture_var.get()},"
            f"autozero={self.autozero_var.get()},"
            f"auto_range={self.auto_range_var.get()},"
            f"range={self.range_var.get()}"
        )

    def _current_config_signature(self):
        return (
            self.mode_var.get(),
            self.nplc_var.get(),
            self.aperture_var.get(),
            self.autozero_var.get(),
            self.auto_range_var.get(),
            self.range_var.get(),
        )

    def _read_loop(self) -> None:
        interval = self._safe_float(self.interval_var.get())
        if interval is None or interval <= 0:
            interval = 0.5

        local_rm = None
        local_dmm = None
        try:
            local_rm = pyvisa.ResourceManager("@ivi")
            local_dmm = local_rm.open_resource(RESOURCE)
            local_dmm.timeout = 5000
            self.rm = local_rm
            self.dmm = local_dmm
            idn = self.dmm.query("*IDN?").strip()
            self._configure_measurement_mode()
            self._last_config_signature = self._current_config_signature()
            self.status_var.set(f"Connected: {idn}")

            while self.running:
                # Live auto-reconfigure when mode/settings change from the UI.
                current_signature = self._current_config_signature()
                if (not self._configured) or (current_signature != self._last_config_signature):
                    self._configure_measurement_mode()
                    self._last_config_signature = current_signature
                    self.dmm.clear()
                    self.root.after(0, self.status_var.set, f"Reconfigured: {self.mode_var.get()}")

                for attempt in range(MAX_RETRIES):
                    try:
                        raw = self.dmm.query("READ?")
                        value = self._parse_reading(raw)
                        now = datetime.now().strftime("%Y-%m-%d %H:%M:%S.%f")[:-3]
                        settings = self._settings_string()
                        self.readings.append((self.mode_var.get(), now, value, settings))

                        if self._start_time is None:
                            self._start_time = time.time()
                        self.chart_times.append(time.time() - self._start_time)
                        self.chart_values.append(0.0 if value is None else value)

                        self.root.after(0, self.value_var.set, self._format_value(value))
                        self.root.after(0, self._apply_alarm_style, value)
                        self.root.after(0, self._update_stats)
                        break
                    except (pyvisa.errors.VisaIOError, ValueError):
                        self.dmm.clear()
                        time.sleep(0.4)
                        if attempt == MAX_RETRIES - 1:
                            raise
                time.sleep(interval)
        except Exception as exc:
            self.root.after(0, self.value_var.set, "ERROR")
            self.root.after(0, self.status_var.set, f"Error: {exc}")
            self.root.after(0, self.value_label.config, {"fg": "#ff4444"})
            self.root.after(0, self.start_btn.config, {"state": "normal"})
            self.root.after(0, self.stop_btn.config, {"state": "disabled"})
        finally:
            for obj in (local_dmm, local_rm):
                if obj:
                    try:
                        obj.close()
                    except Exception:
                        pass
            # Only clear the shared references if they still point at the
            # session this thread owned. Otherwise a newer thread that has
            # already taken over would lose its session here.
            if self.dmm is local_dmm:
                self.dmm = None
            if self.rm is local_rm:
                self.rm = None

    def export_csv(self) -> None:
        if not self.readings:
            messagebox.showwarning("No data", "No readings to export yet.")
            return
        path = filedialog.asksaveasfilename(
            defaultextension=".csv",
            filetypes=[("CSV files", "*.csv"), ("All files", "*.*")],
            initialfile=f"meter_log_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv",
        )
        if not path:
            return
        with open(path, "w", newline="") as f:
            writer = csv.writer(f)
            writer.writerow(["Mode", "Timestamp", "Reading", "Settings"])
            writer.writerows(self.readings)
        messagebox.showinfo("Exported", f"CSV saved:\n{path}")

    def export_excel(self) -> None:
        if not self.readings:
            messagebox.showwarning("No data", "No readings to export yet.")
            return
        try:
            import openpyxl
            from openpyxl.chart import LineChart, Reference
            from openpyxl.styles import Font, PatternFill
        except ImportError:
            messagebox.showerror("Missing package", "openpyxl is not installed.")
            return

        path = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")],
            initialfile=f"meter_log_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx",
        )
        if not path:
            return

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Readings"

        ws.append(["Mode", "Timestamp", "Reading", "Settings"])
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="1F4E79")
            cell.font = Font(bold=True, color="FFFFFF")
        for row in self.readings:
            ws.append(list(row))

        ws.column_dimensions["A"].width = 16
        ws.column_dimensions["B"].width = 26
        ws.column_dimensions["C"].width = 18
        ws.column_dimensions["D"].width = 48

        numeric_rows = [idx + 2 for idx, row in enumerate(self.readings) if row[2] is not None]
        if numeric_rows:
            chart_sheet = wb.create_sheet("Chart")
            chart = LineChart()
            chart.title = "Measurement Trend"
            chart.style = 10
            chart.y_axis.title = "Reading"
            chart.x_axis.title = "Sample #"
            chart.width = 22
            chart.height = 12
            start = min(numeric_rows)
            end = max(numeric_rows)
            data_ref = Reference(ws, min_col=3, min_row=start, max_row=end)
            chart.add_data(data_ref)
            chart_sheet.add_chart(chart, "A1")

        wb.save(path)
        messagebox.showinfo("Exported", f"Excel saved:\n{path}")

    def clear_data(self, reset_display=True) -> None:
        self.readings.clear()
        self.chart_times.clear()
        self.chart_values.clear()
        self._start_time = None
        if reset_display:
            self.value_var.set(f"--- {MODE_CONFIG[self.mode_var.get()]['unit']}")
        self._update_stats()
        self.line.set_data([], [])
        self.canvas.draw_idle()

    def on_close(self) -> None:
        self.running = False
        self.root.destroy()


if __name__ == "__main__":
    root = tk.Tk()
    app = VoltmeterApp(root)
    root.mainloop()
